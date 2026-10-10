"""Build the mini-project report (Markdown source, DOCX, PDF) from report_template.md + facts.json.

Requires a tool environment with: pypandoc_binary, python-docx, weasyprint, matplotlib, lxml.
Fonts for the PDF: Liberation Serif TTFs in the directory given by --fonts (falls back to DejaVu Serif).

    <toolvenv>/bin/python build_report.py --fonts <dir-with-LiberationSerif-*.ttf>

Outputs (in results/final_submission/):
    AI_Image_Inpainting_Mini_Project_Report.md   - Markdown source (figures referenced from figures/)
    AI_Image_Inpainting_Mini_Project_Report.docx - editable Word document (equations as native Word math)
    AI_Image_Inpainting_Mini_Project_Report.pdf  - print-ready PDF
"""
from __future__ import annotations

import argparse
import copy
import html
import io
import json
import re
import shutil
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import pypandoc  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402,F401
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402
from lxml import etree  # noqa: E402
from matplotlib import rcParams  # noqa: E402
from matplotlib.mathtext import math_to_image  # noqa: E402

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
NAME = "AI_Image_Inpainting_Mini_Project_Report"
FACTS = json.loads((HERE / "facts.json").read_text())
TITLE = "Image Inpainting Using U-Net and Gated-Convolution U-Net: A Comparative Study"

METHOD_LABEL = {"gated_unet": "Gated U-Net", "unet": "U-Net", "telea": "OpenCV Telea", "mean_fill": "Mean fill"}
ORDER = ["gated_unet", "unet", "telea", "mean_fill"]

# ----------------------------------------------------------------------------------------------
# References (each verified against Crossref / arXiv / the live web page before inclusion)
# ----------------------------------------------------------------------------------------------
BIB = {
    "bertalmio": 'M. Bertalmio, G. Sapiro, V. Caselles, and C. Ballester, "Image inpainting," in *Proc. 27th Annu. Conf. '
                 'Computer Graphics and Interactive Techniques (SIGGRAPH \'00)*, 2000, pp. 417–424, doi: 10.1145/344779.344972.',
    "telea": 'A. Telea, "An image inpainting technique based on the fast marching method," *Journal of Graphics Tools*, '
             'vol. 9, no. 1, pp. 23–34, 2004, doi: 10.1080/10867651.2004.10487596.',
    "criminisi": 'A. Criminisi, P. Pérez, and K. Toyama, "Region filling and object removal by exemplar-based image '
                 'inpainting," *IEEE Transactions on Image Processing*, vol. 13, no. 9, pp. 1200–1212, 2004, '
                 'doi: 10.1109/TIP.2004.833105.',
    "pathak": 'D. Pathak, P. Krähenbühl, J. Donahue, T. Darrell, and A. A. Efros, "Context encoders: Feature learning by '
              'inpainting," in *Proc. IEEE Conf. Computer Vision and Pattern Recognition (CVPR)*, 2016, pp. 2536–2544, '
              'doi: 10.1109/CVPR.2016.278.',
    "iizuka": 'S. Iizuka, E. Simo-Serra, and H. Ishikawa, "Globally and locally consistent image completion," *ACM '
              'Transactions on Graphics*, vol. 36, no. 4, pp. 1–14, 2017, doi: 10.1145/3072959.3073659.',
    "yu2018": 'J. Yu, Z. Lin, J. Yang, X. Shen, X. Lu, and T. S. Huang, "Generative image inpainting with contextual '
              'attention," in *Proc. IEEE/CVF Conf. Computer Vision and Pattern Recognition (CVPR)*, 2018, '
              'pp. 5505–5514, doi: 10.1109/CVPR.2018.00577.',
    "liu2018": 'G. Liu, F. A. Reda, K. J. Shih, T.-C. Wang, A. Tao, and B. Catanzaro, "Image inpainting for irregular '
               'holes using partial convolutions," in *Computer Vision – ECCV 2018*, Lecture Notes in Computer Science. '
               'Cham, Switzerland: Springer, 2018, pp. 89–105, doi: 10.1007/978-3-030-01252-6_6.',
    "yu2019": 'J. Yu, Z. Lin, J. Yang, X. Shen, X. Lu, and T. Huang, "Free-form image inpainting with gated convolution," '
              'in *Proc. IEEE/CVF Int. Conf. Computer Vision (ICCV)*, 2019, pp. 4470–4479, '
              'doi: 10.1109/ICCV.2019.00457.',
    "unet": 'O. Ronneberger, P. Fischer, and T. Brox, "U-Net: Convolutional networks for biomedical image segmentation," '
            'in *Medical Image Computing and Computer-Assisted Intervention – MICCAI 2015*, Lecture Notes in Computer '
            'Science. Cham, Switzerland: Springer, 2015, pp. 234–241, doi: 10.1007/978-3-319-24574-4_28.',
    "lpips": 'R. Zhang, P. Isola, A. A. Efros, E. Shechtman, and O. Wang, "The unreasonable effectiveness of deep '
             'features as a perceptual metric," in *Proc. IEEE/CVF Conf. Computer Vision and Pattern Recognition '
             '(CVPR)*, 2018, pp. 586–595, doi: 10.1109/CVPR.2018.00068.',
    "lama": 'R. Suvorov, E. Logacheva, A. Mashikhin, A. Remizova, A. Ashukha, A. Silvestrov, N. Kong, H. Goka, K. Park, '
            'and V. Lempitsky, "Resolution-robust large mask inpainting with Fourier convolutions," in *Proc. IEEE/CVF '
            'Winter Conf. Applications of Computer Vision (WACV)*, 2022, pp. 3172–3182, '
            'doi: 10.1109/WACV51458.2022.00323.',
    "repaint": 'A. Lugmayr, M. Danelljan, A. Romero, F. Yu, R. Timofte, and L. Van Gool, "RePaint: Inpainting using '
               'denoising diffusion probabilistic models," in *Proc. IEEE/CVF Conf. Computer Vision and Pattern '
               'Recognition (CVPR)*, 2022, pp. 11451–11461, doi: 10.1109/CVPR52688.2022.01117.',
    "hinton": 'G. E. Hinton and R. R. Salakhutdinov, "Reducing the dimensionality of data with neural networks," '
              '*Science*, vol. 313, no. 5786, pp. 504–507, 2006, doi: 10.1126/science.1127647.',
    "dilated": 'F. Yu and V. Koltun, "Multi-scale context aggregation by dilated convolutions," in *Proc. Int. Conf. '
               'Learning Representations (ICLR)*, 2016, arXiv: 1511.07122.',
    "pytorch": 'A. Paszke *et al.*, "PyTorch: An imperative style, high-performance deep learning library," in '
               '*Advances in Neural Information Processing Systems 32 (NeurIPS)*, 2019, arXiv: 1912.01703.',
    "opencv": 'OpenCV team, "opencv-python-headless" (Python package of the OpenCV library), Python Package Index. '
              '[Online]. Available: https://pypi.org/project/opencv-python-headless/',
    "skimage": 'S. van der Walt, J. L. Schönberger, J. Nunez-Iglesias, F. Boulogne, J. D. Warner, N. Yager, '
               'E. Gouillart, and T. Yu, "scikit-image: Image processing in Python," *PeerJ*, vol. 2, p. e453, 2014, '
               'doi: 10.7717/peerj.453.',
    "tvvgg": 'PyTorch Contributors, "vgg16 — Torchvision documentation." [Online]. Available: '
             'https://pytorch.org/vision/stable/models/generated/torchvision.models.vgg16.html',
    "places": 'B. Zhou, A. Lapedriza, A. Khosla, A. Oliva, and A. Torralba, "Places: A 10 million image database for '
              'scene recognition," *IEEE Transactions on Pattern Analysis and Machine Intelligence*, vol. 40, no. 6, '
              'pp. 1452–1464, 2018, doi: 10.1109/TPAMI.2017.2723009.',
    "placesweb": 'MIT CSAIL, "Places2 download page." [Online]. Available: http://places2.csail.mit.edu/download.html',
    "elu": 'D.-A. Clevert, T. Unterthiner, and S. Hochreiter, "Fast and accurate deep network learning by exponential '
           'linear units (ELUs)," in *Proc. Int. Conf. Learning Representations (ICLR)*, 2016, arXiv: 1511.07289.',
    "odena": 'A. Odena, V. Dumoulin, and C. Olah, "Deconvolution and checkerboard artifacts," *Distill*, vol. 1, no. 10, '
             '2016, doi: 10.23915/distill.00003.',
    "amp": 'P. Micikevicius, S. Narang, J. Alben, G. Diamos, E. Elsen, D. Garcia, B. Ginsburg, M. Houston, O. Kuchaiev, '
           'G. Venkatesh, and H. Wu, "Mixed precision training," in *Proc. Int. Conf. Learning Representations '
           '(ICLR)*, 2018, arXiv: 1710.03740.',
    "torchamp": 'PyTorch Contributors, "Automatic mixed precision package — torch.amp," PyTorch documentation. '
                '[Online]. Available: https://pytorch.org/docs/stable/amp.html',
    "johnson": 'J. Johnson, A. Alahi, and L. Fei-Fei, "Perceptual losses for real-time style transfer and '
               'super-resolution," in *Computer Vision – ECCV 2016*, Lecture Notes in Computer Science. Cham, '
               'Switzerland: Springer, 2016, pp. 694–711, doi: 10.1007/978-3-319-46475-6_43.',
    "vgg": 'K. Simonyan and A. Zisserman, "Very deep convolutional networks for large-scale image recognition," '
           'arXiv: 1409.1556, 2014 (presented at ICLR 2015).',
    "adam": 'D. P. Kingma and J. Ba, "Adam: A method for stochastic optimization," in *Proc. Int. Conf. Learning '
            'Representations (ICLR)*, 2015, arXiv: 1412.6980.',
    "sgdr": 'I. Loshchilov and F. Hutter, "SGDR: Stochastic gradient descent with warm restarts," in *Proc. Int. Conf. '
            'Learning Representations (ICLR)*, 2017, arXiv: 1608.03983.',
    "psnr": 'Q. Huynh-Thu and M. Ghanbari, "Scope of validity of PSNR in image/video quality assessment," *Electronics '
            'Letters*, vol. 44, no. 13, pp. 800–801, 2008, doi: 10.1049/el:20080522.',
    "ssim": 'Z. Wang, A. C. Bovik, H. R. Sheikh, and E. P. Simoncelli, "Image quality assessment: From error visibility '
            'to structural similarity," *IEEE Transactions on Image Processing*, vol. 13, no. 4, pp. 600–612, 2004, '
            'doi: 10.1109/TIP.2003.819861.',
}


# ----------------------------------------------------------------------------------------------
# Placeholder filling
# ----------------------------------------------------------------------------------------------
def lookup(path: str):
    node = FACTS
    for part in path.split("."):
        if not isinstance(node, dict) or part not in node:
            raise KeyError(f"facts.json has no value for '{path}' (failed at '{part}')")
        node = node[part]
    return node


def fmt(value, spec: str | None) -> str:
    if spec is None:
        return str(value)
    if spec == "int":
        return f"{int(value):,}"
    if spec == "pct":
        return f"{value * 100:.1f}"
    if spec == "min":
        return f"{value / 60:.1f}"
    if spec == "hours":
        return f"{value / 3600:.1f}"
    if spec == "gib":
        return f"{value / 1024:.2f}"
    out = format(value, spec)
    return out.replace("-", "−") if out.startswith("-") else out


def mean_sd(block, key, digits):
    return f"{block[key]['mean']:.{digits}f} ± {block[key]['std']:.{digits}f}"


def table_overall(split: str, number: int) -> str:
    o = FACTS[split]["overall"]
    n = FACTS[split]["protocol"]["num_images"]
    title = ("validation set" if split == "val" else "held-out test set") + f" ({n:,} images, 128 × 128)"
    rows = [f"Table: Table {number}. Results on the {title}. Mean ± standard deviation over images; "
            "best value in each column in bold.", "",
            "| Method | PSNR (dB) ↑ | SSIM ↑ | Hole PSNR (dB) ↑ |",
            "|---|---:|---:|---:|"]
    best = {k: max(o[m][k]["mean"] for m in ORDER) for k in ("full_psnr", "full_ssim", "hole_psnr")}
    best.update({k: min(o[m][k]["mean"] for m in ORDER) for k in ("full_mse", "hole_mse")})
    for m in ORDER:
        cells = []
        for k, d in (("full_psnr", 2), ("full_ssim", 4), ("hole_psnr", 2)):
            c = mean_sd(o[m], k, d)
            cells.append(f"**{c}**" if o[m][k]["mean"] == best[k] else c)
        rows.append(f"| {METHOD_LABEL[m]} | " + " | ".join(cells) + " |")
    return "\n".join(rows)


def table_val_vs_test(number: int) -> str:
    rows = [f"Table: Table {number}. Validation versus held-out test results for the two networks (means over images).",
            "",
            "| Network | Split | Images | PSNR (dB) | SSIM | Hole PSNR (dB) |",
            "|---|---|---:|---:|---:|---:|"]
    for m in ("gated_unet", "unet"):
        for split, lab in (("val", "Validation"), ("test", "Test")):
            o = FACTS[split]["overall"][m]
            rows.append(f"| {METHOD_LABEL[m]} | {lab} | {FACTS[split]['protocol']['num_images']:,} | "
                        f"{o['full_psnr']['mean']:.2f} | {o['full_ssim']['mean']:.4f} | {o['hole_psnr']['mean']:.2f} |")
    return "\n".join(rows)


def table_paired(number: int) -> str:
    p = FACTS["test"]["paired"]
    rows = [f"Table: Table {number}. Paired per-image differences on the test set (first method minus second; "
            "3,500 images). CI: 95 % confidence interval of the mean difference over test images.", "",
            "| Comparison | Mean Δ hole PSNR (dB) | 95 % CI (dB) | Mean Δ SSIM | Better on |",
            "|---|---:|---:|---:|---:|"]
    for key in ("gated_unet-unet", "gated_unet-telea", "unet-telea", "telea-mean_fill"):
        a, b = key.split("-")
        h, s = p[key]["hole_psnr"], p[key]["full_ssim"]
        rows.append(f"| {METHOD_LABEL[a]} − {METHOD_LABEL[b]} | {fmt(h['mean'], '+.3f')} | "
                    f"[{fmt(h['ci_lo'], '+.3f')}, {fmt(h['ci_hi'], '+.3f')}] | {fmt(s['mean'], '+.4f')} | "
                    f"{h['frac_better'] * 100:.1f} % |")
    return "\n".join(rows)


def table_groups(group: str, number: int, caption: str, label: dict | None = None) -> str:
    g = FACTS["test"][group]
    keys = list(g["unet"].keys())
    rows = [f"Table: Table {number}. {caption}", "",
            "| " + ("Mask family" if group == "by_mask_family" else "Hole ratio") + " | Images | " +
            " | ".join(METHOD_LABEL[m] for m in ORDER) + " |",
            "|---|---:|" + "---:|" * len(ORDER)]
    for k in keys:
        rows.append(f"| {(label or {}).get(k, k)} | {g['unet'][k]['n']:,} | " +
                    " | ".join(f"{g[m][k]['hole_psnr']['mean']:.2f}" for m in ORDER) + " |")
    return "\n".join(rows)


TABLES = {
    "val_overall": lambda: table_overall("val", 8),
    "test_overall": lambda: table_overall("test", 9),
    "val_vs_test": lambda: table_val_vs_test(10),
    "paired": lambda: table_paired(11),
    "by_family": lambda: table_groups("by_mask_family", 12, "Mean hole PSNR (dB) on the test set by mask family.",
                                      {"rect": "Rectangles", "free_form": "Free-form strokes", "both": "Both"}),
    "by_bucket": lambda: table_groups("by_hole_bucket", 13, "Mean hole PSNR (dB) on the test set by hole ratio."),
}


def fill(text: str) -> str:
    text = re.sub(r"\{\{TABLE:(\w+)\}\}", lambda m: TABLES[m.group(1)](), text)

    def repl(m):
        key, _, spec = m.group(1).partition("|")
        return fmt(lookup(key), spec or None)
    return re.sub(r"\{\{(?!(?:TITLEPAGE|PAGEBREAK|TOC|REFERENCES)\}\})([a-zA-Z0-9_.%\-]+(?:\|[^}]+)?)\}\}", repl, text)


def cite(text: str) -> tuple[str, list[str]]:
    order: list[str] = []

    def repl(m):
        keys = [k.strip().lstrip("@") for k in m.group(1).split(";")]
        for k in keys:
            if k not in BIB:
                raise KeyError(f"citation key '{k}' is not in the bibliography")
            if k not in order:
                order.append(k)
        return ", ".join(f"[{order.index(k) + 1}]" for k in keys)
    return re.sub(r"\[(@[\w]+(?:;\s*@[\w]+)*)\]", repl, text), order


# ----------------------------------------------------------------------------------------------
# Target-specific blocks
# ----------------------------------------------------------------------------------------------
TITLE_LINES = [
    ("A Mini-Project Report", "sub"),
    ("on", "small"),
    (TITLE, "title"),
    ("Submitted in partial fulfilment of the requirements for the award of the degree of", "small"),
    ("[Degree Name] in [Branch / Specialisation]", "sub"),
    ("Submitted by", "small"),
    ("[Student Name]", "sub"),
    ("[Register Number]", "small"),
    ("Under the guidance of", "small"),
    ("[Faculty Guide], [Designation]", "sub"),
    ("[Department]", "sub"),
    ("[Institution]", "sub"),
    ("[Academic Year]", "small"),
]


def titlepage_md() -> str:
    return "\n\n".join((f"**{t}**" if k == "title" else t) for t, k in TITLE_LINES)


def titlepage_openxml() -> str:
    size = {"title": 36, "sub": 28, "small": 24}
    before = {"title": 480, "sub": 120, "small": 360}
    paras = []
    for t, k in TITLE_LINES:
        bold = "<w:b/>" if k in ("title", "sub") else ""
        paras.append(f'<w:p><w:pPr><w:spacing w:before="{before[k]}" w:after="120" w:line="300" '
                     f'w:lineRule="auto"/><w:jc w:val="center"/></w:pPr>'
                     f'<w:r><w:rPr>{bold}<w:sz w:val="{size[k]}"/></w:rPr><w:t xml:space="preserve">'
                     f'{html.escape(t)}</w:t></w:r></w:p>')
    return "```{=openxml}\n" + "".join(paras) + "\n```"


def titlepage_html() -> str:
    lines = "".join(f'<p class="tp-{k}">' + html.escape(t).replace("U-Net", '<span class="nowrap">U-Net</span>') + "</p>"
                    for t, k in TITLE_LINES)
    return f'```{{=html}}\n<section class="title-page">{lines}</section>\n```'


PAGEBREAK = {"docx": "",  # Heading 1 already starts a new page in the DOCX styles
             "html": '```{=html}\n<div class="pagebreak"></div>\n```', "md": ""}


def toc_entries(md: str) -> list[tuple[int, str]]:
    after = md.split("{{TOC}}", 1)[1]
    return [(len(h), t.strip()) for h, t in re.findall(r"^(#{1,2}) (.+)$", after, re.M)]


def toc_md(entries) -> str:
    return "\n".join(("" if lvl == 1 else "    ") + f"- {t}" for lvl, t in entries)


# ----------------------------------------------------------------------------------------------
# Equations for the HTML/PDF route (mathtext SVG, text as paths)
# ----------------------------------------------------------------------------------------------
def render_equations(md: str, eq_dir: Path) -> str:
    eq_dir.mkdir(parents=True, exist_ok=True)
    rcParams["mathtext.fontset"] = "stix"
    rcParams["svg.fonttype"] = "path"
    rcParams["font.size"] = 13.5
    count = 0

    def repl(m):
        nonlocal count
        count += 1
        body = m.group(1).strip().replace("\\frac", "\\dfrac")
        path = eq_dir / f"eq{count:02d}.svg"
        math_to_image(f"${body}$", str(path), dpi=72, format="svg")
        return f'```{{=html}}\n<div class="equation"><img src="{path.as_uri()}" alt="equation {count}"/></div>\n```'
    return re.sub(r"\$\$(.+?)\$\$", repl, md, flags=re.S)


# ----------------------------------------------------------------------------------------------
# DOCX post-processing
# ----------------------------------------------------------------------------------------------
def insert_before_any(parent, new, later_tags):
    """Insert `new` before the first existing child whose tag is in `later_tags` (schema order), else append."""
    later = {qn(t) for t in later_tags}
    for i, child in enumerate(parent):
        if child.tag in later:
            parent.insert(i, new)
            return
    parent.append(new)


TBLPR_AFTER_JC = ["w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                  "w:tblCaption", "w:tblDescription"]
TBLPR_AFTER_BORDERS = ["w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"]
TCPR_AFTER_BORDERS = ["w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
                      "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange"]
SETTINGS_AFTER_UPDATEFIELDS = ["w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids",
                               "m:mathPr", "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping",
                               "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade",
                               "w:captions", "w:readModeInkLockDown", "w:smartTagType", "w:shapeDefaults",
                               "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator"]


def set_font(style, name="Times New Roman", size=None, bold=None, italic=None, color=None):
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = parse_xml(f"<w:rFonts {nsdecls('w')}/>")
        rpr.insert(0, fonts)
    for att in list(fonts.attrib):
        if "Theme" in att:
            del fonts.attrib[att]
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(att), name)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = RGBColor.from_string(color)


def add_page_number_footer(section):
    section.different_first_page_header_footer = True
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for el in (f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>',
               f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>',
               f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>',
               f'<w:r {nsdecls("w")}><w:t>1</w:t></w:r>',
               f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'):
        p._p.append(parse_xml(el))


def toc_sdt(entries) -> str:
    items = "".join(
        f'<w:p><w:pPr><w:pStyle w:val="{"TOC1" if lvl == 1 else "TOC2"}"/></w:pPr>'
        f'<w:r><w:t xml:space="preserve">{html.escape(t)}</w:t></w:r></w:p>' for lvl, t in entries)
    first = ('<w:p><w:pPr><w:pStyle w:val="TOC1"/></w:pPr>'
             '<w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
             '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText></w:r>'
             '<w:r><w:fldChar w:fldCharType="separate"/></w:r></w:p>')
    last = '<w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
    return (f'<w:sdt {nsdecls("w")}><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/>'
            f'<w:docPartUnique/></w:docPartObj></w:sdtPr><w:sdtContent>{first}{items}{last}</w:sdtContent></w:sdt>')


def ensure_toc_styles(doc):
    styles = doc.styles
    for name, indent in (("toc 1", 0), ("toc 2", 0.6)):
        try:
            st = styles[name]
        except KeyError:
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles["Normal"]
        set_font(st, size=12, bold=(name == "toc 1"))
        pf = st.paragraph_format
        pf.left_indent = Cm(indent)
        pf.space_before = Pt(3 if name == "toc 1" else 0)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.15
        pf.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        st.element.set(qn("w:styleId"), "TOC1" if name == "toc 1" else "TOC2")


def postprocess_docx(path: Path, entries):
    doc = Document(str(path))
    styles = doc.styles

    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        if name in [s.name for s in styles]:
            set_font(styles[name], size=12)
    for name in ("Body Text", "First Paragraph"):
        pf = styles[name].paragraph_format
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    styles["Compact"].paragraph_format.line_spacing = 1.15
    styles["Compact"].paragraph_format.space_after = Pt(2)
    for lvl, size in ((1, 16), (2, 13), (3, 12)):
        st = styles[f"Heading {lvl}"]
        set_font(st, size=size, bold=True, italic=False, color="000000")
        pf = st.paragraph_format
        pf.space_before = Pt(18 if lvl == 1 else 12)
        pf.space_after = Pt(8 if lvl == 1 else 6)
        pf.keep_with_next = True
        pf.page_break_before = (lvl == 1)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for name, kw in (("Image Caption", dict(size=10, italic=True)), ("Table Caption", dict(size=10, bold=True)),
                     ("Captioned Figure", {}), ("Figure", {})):
        if name in [s.name for s in styles]:
            set_font(styles[name], **kw)
            pf = styles[name].paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.line_spacing = 1.15
            pf.space_after = Pt(10 if name == "Image Caption" else 4)
            if name in ("Table Caption", "Captioned Figure", "Figure"):
                pf.keep_with_next = True
    ensure_toc_styles(doc)

    # Page setup: A4, 2.54 cm margins, page numbers in the footer (none on the title page).
    for sec in doc.sections:
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.54)
        sec.header_distance = sec.footer_distance = Cm(1.25)
        sec.gutter = Cm(0)
        add_page_number_footer(sec)

    # Tables: horizontal rules (booktabs-like), 10.5 pt text, header row repeated, centred.
    for tbl in doc.tables:
        tpr = tbl._tbl.tblPr
        for old in tpr.findall(qn("w:tblBorders")):
            tpr.remove(old)
        insert_before_any(tpr, parse_xml(
            f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="10" w:color="000000"/>'
            f'<w:bottom w:val="single" w:sz="10" w:color="000000"/>'
            f'<w:insideH w:val="single" w:sz="2" w:color="A6A6A6"/></w:tblBorders>'), TBLPR_AFTER_BORDERS)
        for old in tpr.findall(qn("w:jc")):
            tpr.remove(old)
        insert_before_any(tpr, parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'), TBLPR_AFTER_JC)
        for i, row in enumerate(tbl.rows):
            trpr = row._tr.get_or_add_trPr()  # CT_TrPr children form an xsd:choice, so order is free
            for old in trpr.findall(qn("w:cantSplit")) + trpr.findall(qn("w:tblHeader")):
                trpr.remove(old)
            trpr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            if i == 0:
                trpr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            for tc in row._tr.findall(qn("w:tc")):
                for p in tc.iter(qn("w:p")):
                    for r in p.iter(qn("w:r")):
                        rpr = r.get_or_add_rPr()
                        rpr.get_or_add_sz().set(qn("w:val"), "21")
                        if i == 0:
                            rpr.get_or_add_b()
                if i == 0:
                    tcpr = tc.get_or_add_tcPr()
                    insert_before_any(tcpr, parse_xml(
                        f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:color="000000"/>'
                        f'</w:tcBorders>'), TCPR_AFTER_BORDERS)

    # Table of contents field at the placeholder paragraph.
    for p in doc.paragraphs:
        if p.text.strip() == "TOC_PLACEHOLDER":
            p._p.addnext(parse_xml(toc_sdt(entries)))
            p._p.getparent().remove(p._p)
            break
    else:
        raise RuntimeError("TOC placeholder not found in DOCX")

    # Ask Word to refresh fields (fills TOC page numbers) when the document is opened.
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        insert_before_any(settings, parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>'),
                          SETTINGS_AFTER_UPDATEFIELDS)

    # Reference list: hanging indent, single-ish spacing, left aligned.
    in_refs = False
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading 1"):
            in_refs = p.text.strip() == "References"
            continue
        if in_refs and p.text.startswith("["):
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Cm(1.0), Cm(-1.0)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.line_spacing = 1.15
            pf.space_after = Pt(6)

    cp = doc.core_properties
    cp.title, cp.subject = TITLE, "Mini-project report"
    cp.author, cp.keywords = "[Student Name]", "image inpainting; U-Net; gated convolution"
    doc.save(str(path))


# ----------------------------------------------------------------------------------------------
# HTML / PDF
# ----------------------------------------------------------------------------------------------
CSS = """
@font-face { font-family: 'RepSerif'; src: url('{F}/LiberationSerif-Regular.ttf'); }
@font-face { font-family: 'RepSerif'; font-weight: bold; src: url('{F}/LiberationSerif-Bold.ttf'); }
@font-face { font-family: 'RepSerif'; font-style: italic; src: url('{F}/LiberationSerif-Italic.ttf'); }
@font-face { font-family: 'RepSerif'; font-weight: bold; font-style: italic; src: url('{F}/LiberationSerif-BoldItalic.ttf'); }
@page { size: A4; margin: 25.4mm 25.4mm 25.4mm 25.4mm;
        @bottom-center { content: counter(page); font-family: 'RepSerif', 'DejaVu Serif'; font-size: 10pt; } }
@page :first { @bottom-center { content: none; } }
html { font-family: 'RepSerif', 'DejaVu Serif', serif; font-size: 12pt; line-height: 1.5; color: #000; }
body { margin: 0; }
p { text-align: justify; hyphens: auto; margin: 0 0 6pt 0; }
h1 { font-size: 16pt; margin: 0 0 10pt 0; page-break-before: always; break-after: avoid; line-height: 1.25; }
h2 { font-size: 13pt; margin: 14pt 0 6pt 0; break-after: avoid; line-height: 1.25; }
h3 { font-size: 12pt; margin: 12pt 0 4pt 0; break-after: avoid; }
h2#keywords { margin-top: 18pt; }
.title-page { page-break-after: always; text-align: center; padding-top: 25mm; }
.title-page p { text-align: center; margin: 0; line-height: 1.35; }
.tp-title { font-size: 18pt; font-weight: bold; margin: 22pt 0 26pt 0 !important; line-height: 1.3; }
.tp-sub { font-size: 14pt; font-weight: bold; margin-top: 8pt !important; }
.tp-small { font-size: 12pt; margin-top: 20pt !important; }
.pagebreak { page-break-after: always; }
figure { margin: 10pt 0 12pt 0; text-align: center; page-break-inside: avoid; }
figure img { max-width: 100%; }
figcaption { font-size: 10pt; font-style: italic; line-height: 1.3; margin-top: 4pt; text-align: center; }
table { border-collapse: collapse; margin: 6pt auto 14pt auto; font-size: 10pt; line-height: 1.3;
        page-break-inside: avoid; border-top: 1.2pt solid #000; border-bottom: 1.2pt solid #000; }
caption { caption-side: top; font-size: 10pt; font-weight: bold; padding-bottom: 4pt; line-height: 1.3; }
th { border-bottom: 0.8pt solid #000; padding: 3pt 6pt; text-align: left; vertical-align: bottom; }
td { padding: 2.5pt 6pt; border-top: 0.3pt solid #bbb; vertical-align: top; }
td p, th p { margin: 0; text-align: inherit; }
td:first-child, th:first-child { white-space: nowrap; }
.tblwrap { page-break-inside: avoid; break-inside: avoid; margin: 6pt 0 14pt 0; }
.tblwrap table { margin: 0 auto; }
.nowrap { white-space: nowrap; }
code { font-family: 'DejaVu Sans Mono', monospace; font-size: 9.5pt; }
.equation { text-align: center; margin: 6pt 0 8pt 0; page-break-inside: avoid; }
.equation img { height: auto; }
ol, ul { margin: 0 0 6pt 0; padding-left: 22pt; }
li { margin-bottom: 3pt; text-align: justify; }
li > p { margin: 0; }
nav.toc ul { list-style: none; padding-left: 0; margin: 0; }
nav.toc li { margin: 0; text-align: left; }
nav.toc li.l1 { font-weight: bold; margin-top: 5pt; }
nav.toc li.l2 { padding-left: 18pt; font-weight: normal; }
nav.toc a { color: #000; text-decoration: none; }
nav.toc a::after { content: leader('.') target-counter(attr(href), page); }
.references p { text-align: left; padding-left: 26pt; text-indent: -26pt; margin-bottom: 5pt; line-height: 1.3; }
"""


def build_html(md_html: str, entries, fonts: Path) -> str:
    body = pypandoc.convert_text(md_html, "html5", format="markdown+raw_attribute+implicit_figures",
                                 extra_args=["--wrap=none"])
    tree = etree.fromstring(f"<div>{body}</div>", etree.HTMLParser())
    heads = [h for h in tree.iter("h1", "h2") if h.get("id")]
    texts = {"".join(h.itertext()).strip(): h.get("id") for h in heads}
    lis = []
    for lvl, t in entries:
        hid = texts.get(t)
        if hid is None:
            raise RuntimeError(f"heading id not found for TOC entry '{t}'")
        lis.append(f'<li class="l{lvl}"><a href="#{hid}">{html.escape(t)}</a></li>')
    toc = f'<nav class="toc"><ul>{"".join(lis)}</ul></nav>'
    if "<p>TOC_PLACEHOLDER</p>" not in body:
        raise RuntimeError("TOC placeholder missing in HTML")
    body = body.replace("<p>TOC_PLACEHOLDER</p>", toc)
    body = re.sub(r'(<h1 id="references">References</h1>)(.*)$', r'\1<div class="references">\2</div>', body,
                  flags=re.S)
    body = body.replace("<table", '<div class="tblwrap"><table').replace("</table>", "</table></div>")
    css = CSS.replace("{F}", fonts.as_uri())
    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{html.escape(TITLE)}</title>'
            f'<style>{css}</style></head><body>{body}</body></html>')


# ----------------------------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--fonts", type=Path, required=True)
    args = ap.parse_args()

    template = (HERE / "report_template.md").read_text()
    filled = fill(template)
    filled = re.sub(r"(\d) %", "\\1\u00a0%", filled)  # keep "40 %" on one line
    leftovers = sorted(set(re.findall(r"\{\{(?!TITLEPAGE|PAGEBREAK|TOC|REFERENCES)[^}]*\}\}", filled)))
    if leftovers:
        raise RuntimeError(f"unfilled placeholders: {leftovers}")
    filled, order = cite(filled)
    refs = "\n\n".join(f"[{i}] {BIB[k]}" for i, k in enumerate(order, 1))
    filled = filled.replace("{{REFERENCES}}", refs)
    unused = sorted(set(BIB) - set(order))
    entries = toc_entries(filled)

    # Markdown source
    md_src = (filled.replace("{{TITLEPAGE}}", titlepage_md()).replace("{{PAGEBREAK}}", "")
              .replace("{{TOC}}", toc_md(entries)))
    (OUT / f"{NAME}.md").write_text(md_src)

    # DOCX
    md_docx = (filled.replace("{{TITLEPAGE}}", titlepage_openxml()).replace("{{PAGEBREAK}}", PAGEBREAK["docx"])
               .replace("{{TOC}}", "TOC_PLACEHOLDER"))
    docx_path = OUT / f"{NAME}.docx"
    pypandoc.convert_text(md_docx, "docx", format="markdown+raw_attribute+implicit_figures", outputfile=str(docx_path),
                          extra_args=[f"--resource-path={OUT}", "--wrap=none"])
    postprocess_docx(docx_path, entries)

    # PDF
    with tempfile.TemporaryDirectory() as tmp:
        md_html = (filled.replace("{{TITLEPAGE}}", titlepage_html()).replace("{{PAGEBREAK}}", PAGEBREAK["html"])
                   .replace("{{TOC}}", "TOC_PLACEHOLDER"))
        md_html = render_equations(md_html, Path(tmp) / "eq")
        md_html = md_html.replace("](figures/", f"]({(OUT / 'figures').as_uri()}/")
        page = build_html(md_html, entries, args.fonts)
        (HERE / "report_preview.html").write_text(page)
        from weasyprint import HTML
        HTML(string=page, base_url=str(OUT)).write_pdf(str(OUT / f"{NAME}.pdf"))
        (HERE / "report_preview.html").unlink()

    print("citations used (in order):", len(order), "| bibliography entries not cited:", unused or "none")
    print("TOC entries:", len(entries))
    for ext in ("md", "docx", "pdf"):
        p = OUT / f"{NAME}.{ext}"
        print(f"wrote {p} ({p.stat().st_size / 1024:.0f} KiB)")


if __name__ == "__main__":
    main()
